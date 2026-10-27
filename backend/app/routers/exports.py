import csv
import io
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.db import get_db
from app.deps import RateLimiter, require_user
from app.models import Transaction, TxType, User
from app.services import finance

router = APIRouter(prefix="/api/exports", tags=["exports"])
export_limiter = RateLimiter(limit=20, window_seconds=60)
MAX_EXPORT_ROWS = 20000
HEADERS = ["Date", "Type", "Category", "Description", "Amount (INR)"]


def _csv_safe(v: str | None) -> str:
    """Neutralise spreadsheet formula injection in text cells."""
    v = v or ""
    return "'" + v if v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


def _scope(scope: str, start: date | None, end: date | None) -> tuple[TxType | None, date | None, date | None]:
    if scope == "all":
        return None, None, None
    if scope == "income":
        return TxType.INCOME, None, None
    if scope == "expenses":
        return TxType.EXPENSE, None, None
    if scope in ("this_month", "last_month", "this_year"):
        s, e = finance.resolve_period(scope, None, None)
        return None, s, e
    if scope == "custom":
        s, e = finance.resolve_period("custom", start, end)
        return None, s, e
    raise HTTPException(422, "Unknown export scope")


def _fetch(db: Session, user: User, tx_type, s, e):
    stmt = select(Transaction).where(Transaction.user_id == user.id)
    if tx_type:
        stmt = stmt.where(Transaction.type == tx_type)
    if s:
        stmt = stmt.where(Transaction.transaction_date >= s)
    if e:
        stmt = stmt.where(Transaction.transaction_date <= e)
    stmt = stmt.order_by(Transaction.transaction_date.asc(), Transaction.id.asc()).limit(MAX_EXPORT_ROWS + 1)
    rows = db.scalars(stmt).unique().all()
    if len(rows) > MAX_EXPORT_ROWS:
        raise HTTPException(413, f"Too many records to export at once (limit {MAX_EXPORT_ROWS}). Narrow the date range.")
    return rows


def _summary(rows):
    inc = sum((finance.q2(r.amount) for r in rows if r.type == TxType.INCOME), finance.ZERO)
    exp = sum((finance.q2(r.amount) for r in rows if r.type == TxType.EXPENSE), finance.ZERO)
    return inc, exp, inc - exp


def _label(tx_type, s, e) -> str:
    parts = []
    if tx_type:
        parts.append("Income only" if tx_type == TxType.INCOME else "Expenses only")
    if s and e:
        parts.append(f"{s.isoformat()} to {e.isoformat()}")
    return ", ".join(parts) or "All transactions"


def _csv(rows, label) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADERS)
    for r in rows:
        w.writerow([r.transaction_date.isoformat(), r.type.value, _csv_safe(r.category.name),
                    _csv_safe(r.description), f"{finance.q2(r.amount):.2f}"])
    inc, exp, net = _summary(rows)
    w.writerow([])
    w.writerow(["Total income", "", "", "", f"{inc:.2f}"])
    w.writerow(["Total expenses", "", "", "", f"{exp:.2f}"])
    w.writerow(["Net balance", "", "", "", f"{net:.2f}"])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8 (₹, accents) correctly


def _xlsx(rows, label) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Income & Expense Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Filter: {label}"])
    ws.append([])
    ws.append(HEADERS)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r.transaction_date, r.type.value, _csv_safe(r.category.name), _csv_safe(r.description),
                   float(finance.q2(r.amount))])
    first, last = 5, 4 + len(rows)
    for row in ws.iter_rows(min_row=first, max_row=last):
        row[0].number_format = "yyyy-mm-dd"
        row[4].number_format = '"₹"#,##0.00'
    inc, exp, net = _summary(rows)
    ws.append([])
    for name, val in (("Total income", inc), ("Total expenses", exp), ("Net balance", net)):
        ws.append([name, "", "", "", float(val)])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=5).number_format = '"₹"#,##0.00'
        ws.cell(row=ws.max_row, column=5).font = Font(bold=True)
    for col, width in zip("ABCDE", (14, 11, 22, 44, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _pdf(rows, label, user: User) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
                            title="Financial summary")
    styles = getSampleStyleSheet()
    inc, exp, net = _summary(rows)
    money = lambda v: f"Rs. {v:,.2f}"  # built-in PDF fonts have no rupee glyph
    story = [
        Paragraph("Financial summary", styles["Title"]),
        Paragraph(f"{user.full_name} &middot; Filter: {label} &middot; Generated {date.today().isoformat()}", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]
    summary = Table([["Total income", money(inc)], ["Total expenses", money(exp)], ["Net balance", money(net)]],
                    colWidths=[50 * mm, 45 * mm])
    summary.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
                                 ("ALIGN", (1, 0), (1, -1), "RIGHT")]))
    story += [summary, Spacer(1, 8 * mm)]
    if rows:
        data = [["Date", "Type", "Category", "Description", "Amount"]]
        for r in rows:
            desc = Paragraph((r.description or "").replace("&", "&amp;").replace("<", "&lt;"), styles["BodyText"])
            data.append([r.transaction_date.isoformat(), r.type.value.title(), r.category.name, desc, f"{finance.q2(r.amount):,.2f}"])
        t = Table(data, repeatRows=1, colWidths=[24 * mm, 20 * mm, 36 * mm, 70 * mm, 26 * mm])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                               ("ALIGN", (4, 0), (4, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")])]))
        story.append(t)
    else:
        story.append(Paragraph("No transactions match this filter.", styles["Normal"]))
    doc.build(story)
    return out.getvalue()


@router.get("/transactions", dependencies=[Depends(export_limiter)])
def export_transactions(
    format: str = Query(pattern="^(csv|xlsx|pdf)$"),
    scope: str = "all", start_date: date | None = None, end_date: date | None = None,
    db: Session = Depends(get_db), user: User = Depends(require_user),
):
    tx_type, s, e = _scope(scope, start_date, end_date)
    rows = _fetch(db, user, tx_type, s, e)
    label = _label(tx_type, s, e)
    stamp = date.today().isoformat()
    if format == "csv":
        body, mime = _csv(rows, label), "text/csv; charset=utf-8"
    elif format == "xlsx":
        body, mime = _xlsx(rows, label), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        body, mime = _pdf(rows, label, user), "application/pdf"
    return Response(content=body, media_type=mime,
                    headers={"Content-Disposition": f'attachment; filename="transactions-{stamp}.{format}"',
                             "Cache-Control": "no-store"})
