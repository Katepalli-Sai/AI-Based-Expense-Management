import csv
import io
import json
from unittest.mock import patch

from openpyxl import load_workbook

from app.core.config import get_settings
from app.services import ai
from tests.conftest import add_tx


def seed(client, h, cats):
    add_tx(client, h, "INCOME", "1000", cats["salary"], "2026-10-01", "Pay")
    add_tx(client, h, "EXPENSE", "200", cats["food"], "2026-10-02", "=HYPERLINK(\"x\")")
    add_tx(client, h, "EXPENSE", "50", cats["food"], "2026-09-15", "Old")


def export(client, h, **params):
    return client.get("/api/exports/transactions", headers=h, params=params)


def test_csv_filters_and_summary(client, user_a, cats):
    seed(client, user_a, cats)
    r = export(client, user_a, format="csv", scope="expenses")
    assert r.status_code == 200 and "text/csv" in r.headers["content-type"]
    rows = list(csv.reader(io.StringIO(r.text.lstrip("\ufeff"))))
    assert rows[0] == ["Date", "Type", "Category", "Description", "Amount (INR)"]
    body = [x for x in rows[1:] if x and x[1] == "EXPENSE"]
    assert len(body) == 2 and all(x[1] != "INCOME" for x in rows[1:3])
    assert ["Total expenses", "", "", "", "250.00"] in rows and ["Total income", "", "", "", "0.00"] in rows
    assert "'=HYPERLINK" in r.text  # formula injection neutralised
    c = list(csv.reader(io.StringIO(export(client, user_a, format="csv", scope="custom", start_date="2026-10-01", end_date="2026-10-31").text.lstrip("\ufeff"))))
    assert ["Net balance", "", "", "", "800.00"] in c
    assert len([x for x in c[1:] if x and x[0].startswith("2026-")]) == 2


def test_xlsx_and_pdf_formats(client, user_a, cats):
    seed(client, user_a, cats)
    r = export(client, user_a, format="xlsx", scope="income")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[4]][:2] == ["Date", "Type"]
    assert ws["B5"].value == "INCOME" and ws["E5"].value == 1000
    p = export(client, user_a, format="pdf", scope="all")
    assert p.status_code == 200 and p.content.startswith(b"%PDF") and p.headers["content-type"] == "application/pdf"
    assert export(client, user_a, format="exe").status_code == 422
    assert export(client, user_a, format="csv", scope="bogus").status_code == 422


def test_export_requires_auth_and_isolated(client, user_a, user_b, cats):
    seed(client, user_a, cats)
    assert client.get("/api/exports/transactions", params={"format": "csv"}).status_code == 401
    r = export(client, user_b, format="csv")
    assert "Pay" not in r.text and ["Total income", "", "", "", "0.00"] in list(csv.reader(io.StringIO(r.text.lstrip("\ufeff"))))
    x = load_workbook(io.BytesIO(export(client, user_b, format="xlsx").content)).active
    assert all("Pay" != (c.value or "") for row in x.iter_rows() for c in row)


# ------------------------------------------------------------------ AI
def test_ai_missing_key_does_not_break_app(client, user_a, cats):
    assert not get_settings().gemini_configured
    assert client.get("/api/ai/status", headers=user_a).json() == {"configured": False}
    assert client.post("/api/ai/chat", headers=user_a, json={"message": "hi"}).status_code == 503
    seed(client, user_a, cats)
    assert client.post("/api/ai/insights/refresh", headers=user_a).status_code == 503
    g = client.get("/api/ai/insights", headers=user_a).json()  # still serves exact facts
    assert g["ai_configured"] is False and g["facts"] and g["insights"] is None
    assert client.get("/api/dashboard/summary", headers=user_a).status_code == 200  # non-AI features unaffected


def test_insights_empty_state_never_calls_gemini(client, user_a):
    with patch("app.services.ai.generate_text") as gen:
        r = client.post("/api/ai/insights/refresh", headers=user_a)
        assert r.status_code == 200 and r.json()["enough_data"] is False and r.json()["insights"] is None
        gen.assert_not_called()


FAKE = json.dumps({"summary": "Spending is steady.", "items": [
    {"kind": "observation", "title": "Food", "detail": "Food was your top category."},
    {"kind": "weird", "title": "X", "detail": "Y"}, {"title": "no detail"}]})


def test_insights_generation_cache_and_invalidation(client, user_a, cats):
    seed(client, user_a, cats)
    with patch("app.services.ai.generate_text", return_value="```json\n" + FAKE + "\n```") as gen, \
         patch("app.routers.ai.ai.ai_configured", return_value=True):
        r = client.post("/api/ai/insights/refresh", headers=user_a).json()
        assert r["insights"]["summary"] == "Spending is steady." and len(r["insights"]["items"]) == 2
        assert r["insights"]["items"][1]["kind"] == "observation"  # unknown kind normalised
        client.post("/api/ai/insights/refresh", headers=user_a)
        assert gen.call_count == 1  # unchanged data -> cached, no new Gemini call
        assert client.get("/api/ai/insights", headers=user_a).json()["stale"] is False
        add_tx(client, user_a, "EXPENSE", "5", cats["food"], "2026-10-03")
        assert client.get("/api/ai/insights", headers=user_a).json()["stale"] is True  # invalidated by data change
        assert gen.call_count == 1  # GET never calls Gemini
        client.post("/api/ai/insights/refresh", headers=user_a)
        assert gen.call_count == 2
        # only aggregates are sent: no raw descriptions
        sent = gen.call_args.args[0]
        assert "HYPERLINK" not in sent and "Pay" not in sent.replace("Payment", "")


def test_gemini_errors_are_handled(client, user_a, cats):
    seed(client, user_a, cats)
    for status in (429, 504, 502):
        with patch("app.services.ai.generate_text", side_effect=ai.AIError("friendly", status)), \
             patch("app.routers.ai.ai.ai_configured", return_value=True):
            assert client.post("/api/ai/insights/refresh", headers=user_a).status_code == status
            assert client.post("/api/ai/chat", headers=user_a, json={"message": "hi"}).status_code == status
    with patch("app.services.ai.generate_text", return_value="not json at all"), patch("app.routers.ai.ai.ai_configured", return_value=True):
        assert client.post("/api/ai/insights/refresh", headers=user_a).status_code == 502


def test_chat_conversations_and_isolation(client, user_a, user_b, cats):
    seed(client, user_a, cats)
    with patch("app.services.ai.generate_text", return_value="Facts: you spent 200.") as gen, \
         patch("app.routers.ai.ai.ai_configured", return_value=True):
        r = client.post("/api/ai/chat", headers=user_a, json={"message": "How much did I spend?"}).json()
        cid = r["conversation"]["id"]
        assert r["assistant_message"]["content"].startswith("Facts")
        client.post("/api/ai/chat", headers=user_a, json={"message": "And last month?", "conversation_id": cid})
        assert "How much did I spend?" in gen.call_args.args[0]  # follow-up context
        assert "Pay" not in gen.call_args.args[0].split("CONVERSATION_SO_FAR")[0].replace("Payment", "")  # no raw descriptions
        msgs = client.get(f"/api/ai/conversations/{cid}/messages", headers=user_a).json()["items"]
        assert [m["role"] for m in msgs] == ["user", "assistant", "user", "assistant"]
        assert len(client.get("/api/ai/conversations", headers=user_a).json()["items"]) == 1
        # user B can't read, continue, or delete A's conversation
        assert client.get("/api/ai/conversations", headers=user_b).json()["items"] == []
        assert client.get(f"/api/ai/conversations/{cid}/messages", headers=user_b).status_code == 404
        assert client.post("/api/ai/chat", headers=user_b, json={"message": "x", "conversation_id": cid}).status_code == 404
        assert client.delete(f"/api/ai/conversations/{cid}", headers=user_b).status_code == 404
        assert client.delete(f"/api/ai/conversations/{cid}", headers=user_a).status_code == 204


# ------------------------------------------------------------------ admin
def test_admin_insights_expose_no_individual_financials(client, user_a, cats, admin_headers):
    seed(client, user_a, cats)
    stats = client.get("/api/admin/dashboard", headers=admin_headers).json()
    assert stats["total_users"] == 2 and stats["active_users"] == 2
    dumped = json.dumps(stats) + client.get("/api/admin/users", headers=admin_headers).text
    for secret in ("Pay", "HYPERLINK", "1000", "password_hash", "$argon2"):
        assert secret not in dumped.replace("Payment", "")
    with patch("app.services.ai.generate_text", return_value=FAKE) as gen, patch("app.routers.admin.ai.ai_configured", return_value=True):
        r = client.post("/api/admin/ai/insights/refresh", headers=admin_headers)
        assert r.status_code == 200 and r.json()["insights"]["summary"]
        sent = gen.call_args.args[0]
        assert "PLATFORM_AGGREGATES" in sent and "a@example.com" not in sent and "Pay" not in sent and "HYPERLINK" not in sent
    assert client.post("/api/admin/ai/insights/refresh", headers=admin_headers).status_code == 503  # key missing


def test_admin_user_management_and_last_admin_protection(client, user_a, admin_headers):
    users = client.get("/api/admin/users", headers=admin_headers, params={"search": "alice"}).json()
    assert users["total"] == 1 and "password_hash" not in users["items"][0]
    uid = users["items"][0]["id"]
    assert client.put(f"/api/admin/users/{uid}/status", headers=admin_headers, json={"account_status": "INACTIVE"}).status_code == 200
    assert client.get("/api/auth/me", headers=user_a).status_code == 401  # session ended
    assert client.post("/api/auth/login", json={"email": "a@example.com", "password": "password123"}).status_code == 403
    client.put(f"/api/admin/users/{uid}/status", headers=admin_headers, json={"account_status": "ACTIVE"})
    assert client.post("/api/auth/login", json={"email": "a@example.com", "password": "password123"}).status_code == 200
    admin_id = client.get("/api/auth/me", headers=admin_headers).json()["id"]
    r = client.put(f"/api/admin/users/{admin_id}/status", headers=admin_headers, json={"account_status": "INACTIVE"})
    assert r.status_code == 409


def test_admin_category_management(client, user_a, cats, admin_headers):
    r = client.post("/api/admin/categories", headers=admin_headers, json={"name": "Pets", "type": "EXPENSE", "description": "Pet care"})
    assert r.status_code == 201
    assert client.post("/api/admin/categories", headers=admin_headers, json={"name": "pets", "type": "EXPENSE"}).status_code == 409
    assert client.post("/api/admin/categories", headers=admin_headers, json={"name": "Pets", "type": "INCOME"}).status_code == 201  # same name, other type ok
    cid = r.json()["id"]
    assert client.put(f"/api/admin/categories/{cid}", headers=admin_headers, json={"name": "Pet Care"}).json()["name"] == "Pet Care"
    # history survives deactivation; new use is blocked
    old = add_tx(client, user_a, "EXPENSE", "10", cats["food"], "2026-10-01")
    client.put(f"/api/admin/categories/{cats['food']}", headers=admin_headers, json={"is_active": False})
    assert client.get(f"/api/transactions/{old['id']}", headers=user_a).json()["category_name"] == "Food & Dining"
    assert cats["food"] not in [c["id"] for c in client.get("/api/categories", headers=user_a).json()]
    body = {"type": "EXPENSE", "amount": "5", "category_id": cats["food"], "transaction_date": "2026-10-01"}
    assert client.post("/api/transactions", headers=user_a, json=body).status_code == 422
    assert client.delete(f"/api/admin/categories/{cid}", headers=admin_headers).status_code == 405  # no hard delete
